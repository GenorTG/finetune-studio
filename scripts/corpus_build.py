"""Build the Korvane test corpus: authored sources -> real office/PDF/HTML/CSV files.

    .venv/bin/python scripts/corpus_build.py --lane hr            # one lane
    .venv/bin/python scripts/corpus_build.py --all                # every lane
    .venv/bin/python scripts/corpus_build.py --lane hr --check    # only validate the front matter

Layout (``tests/corpus/korvane/``): ``src/<lane>/*.src`` (authored), ``files/<lane>/`` (built, committed so the corpus is
byte-frozen), ``manifest/<lane>.jsonl`` (ground-truth facts, see ``scripts/corpus_check.py``).

A source file is a front matter block, then a body:

    ---
    out: employee_handbook_2024.pdf      # extension picks the builder (see below)
    tier: core                           # core | extended
    scan: false                          # pdf only: render pages as noisy images (OCR path)
    ---
    # Heading 1 / ## Heading 2 / ### Heading 3
    Plain paragraphs, **bold**, *italic*.
    - bullets            1. numbered items
    | col a | col b |     (pipe tables; first row = header; `|---|---|` separator line optional)
    [[chart bar "Title" | Q1=12.4, Q2=15.1 | unit=EUR m]]      (also: chart line ...) drawn as an image, values labelled
    \\pagebreak

Builders by extension of ``out``:
  docx / pdf / doc / rtf / odt : body -> docx (python-docx) -> LibreOffice for the other formats
  pptx : slides separated by a line ``=== slide``; ``# Title``, bullets, tables, ``Notes:`` block
  xlsx / xls : ``## Sheet name`` then a pipe table per sheet (numbers become numbers; ``=SUM(..)`` stays a formula)
  html / txt / md / csv / eml / json : body is copied verbatim (the author writes the final bytes)
"""
from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1] / "tests" / "corpus" / "korvane"
VERBATIM = {"html", "txt", "md", "csv", "eml", "json"}
FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
CHART_RE = re.compile(r'^\[\[chart\s+(bar|line)\s+"([^"]*)"\s*\|\s*([^|\]]+?)\s*(?:\|\s*([^\]]*))?\]\]\s*$')
TABLE_SEP = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


class SourceError(ValueError):
    pass


def parse_source(path: Path) -> tuple[dict[str, str], str]:
    text = path.read_text(encoding="utf-8")
    m = re.match(r"^---\n(.*?)\n---\n?(.*)$", text, re.DOTALL)
    if not m:
        raise SourceError(f"{path.name}: missing front matter block")
    meta: dict[str, str] = {}
    for line in m.group(1).splitlines():
        if line.strip() and ":" in line:
            k, v = line.split(":", 1)
            meta[k.strip()] = v.split("#")[0].strip() if k.strip() != "title" else v.strip()
    if not meta.get("out") or "." not in meta["out"]:
        raise SourceError(f"{path.name}: front matter needs `out: name.ext`")
    if meta.get("tier", "core") not in ("core", "extended"):
        raise SourceError(f"{path.name}: tier must be core or extended")
    return meta, m.group(2)


# ── charts (PIL only: no matplotlib dependency) ─────────────────────────────────

def render_chart(kind: str, title: str, data: list[tuple[str, float]], unit: str, dest: Path) -> None:
    from PIL import Image, ImageDraw, ImageFont

    w, h, pad = 900, 520, 70
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    f, fb = ImageFont.truetype(FONT, 18), ImageFont.truetype(FONT_BOLD, 22)
    d.text((pad, 18), title + (f"  ({unit})" if unit else ""), font=fb, fill="black")
    top = max(v for _, v in data) * 1.15 or 1
    base, area = h - pad, h - 2 * pad - 30
    d.line([(pad, base), (w - pad, base)], fill="#444", width=2)
    n, step = len(data), (w - 2 * pad) / max(len(data), 1)
    pts = []
    for i, (label, v) in enumerate(data):
        x0 = pad + i * step + step * 0.18
        y = base - v / top * area
        if kind == "bar":
            d.rectangle([x0, y, x0 + step * 0.64, base], fill="#2d6a9f")
            d.text((x0 + 4, y - 24), f"{v:g}", font=f, fill="black")
        else:
            cx = pad + i * step + step / 2
            pts.append((cx, y))
            d.ellipse([cx - 5, y - 5, cx + 5, y + 5], fill="#2d6a9f")
            d.text((cx - 14, y - 28), f"{v:g}", font=f, fill="black")
        d.text((pad + i * step + step * 0.2, base + 8), label, font=f, fill="black")
    if kind == "line" and len(pts) > 1:
        d.line(pts, fill="#2d6a9f", width=3)
    img.save(dest)
    _ = n


def parse_chart(line: str) -> tuple[str, str, list[tuple[str, float]], str] | None:
    m = CHART_RE.match(line.strip())
    if not m:
        return None
    data = []
    for part in m.group(3).split(","):
        label, _, val = part.partition("=")
        data.append((label.strip(), float(val)))
    unit = (m.group(4) or "").replace("unit=", "").strip()
    return m.group(1), m.group(2), data, unit


# ── docx ───────────────────────────────────────────────────────────────────────

def _inline(par, text: str) -> None:
    for chunk in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            par.add_run(chunk[2:-2]).bold = True
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            par.add_run(chunk[1:-1]).italic = True
        elif chunk:
            par.add_run(chunk)


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def add_table(doc, rows: list[list[str]]) -> None:
    cols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=cols)
    t.style = "Table Grid"
    for i, r in enumerate(rows):
        for j in range(cols):
            cell = t.cell(i, j)
            cell.text = r[j] if j < len(r) else ""
            if i == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True


def build_docx(body: str, dest: Path, workdir: Path) -> None:
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    lines, i, chart_n = body.splitlines(), 0, 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
        elif line.strip() == "\\pagebreak":
            doc.add_page_break()
            i += 1
        elif (chart := parse_chart(line)):
            chart_n += 1
            png = workdir / f"chart{chart_n}.png"
            render_chart(chart[0], chart[1], chart[2], chart[3], png)
            doc.add_picture(str(png), width=Inches(6))
            i += 1
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not TABLE_SEP.match(lines[i]):
                    rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            doc.add_paragraph()
        elif (hm := re.match(r"^(#{1,3})\s+(.*)$", line)):
            doc.add_heading(hm.group(2), level=len(hm.group(1)))
            i += 1
        elif re.match(r"^\s*[-*]\s+", line):
            _inline(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*]\s+", "", line))
            i += 1
        elif re.match(r"^\s*\d+[.)]\s+", line):
            _inline(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", line))
            i += 1
        elif line.startswith(">"):
            _inline(doc.add_paragraph(style="Intense Quote"), line.lstrip("> "))
            i += 1
        else:
            para = [line]
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,3}\s|\s*[-*]\s|\s*\d+[.)]\s|\||>|\[\[chart|\\pagebreak)", lines[i]):
                para.append(lines[i].rstrip())
                i += 1
            _inline(doc.add_paragraph(), " ".join(para))
    doc.save(dest)


def soffice(src: Path, fmt: str, outdir: Path) -> Path:
    # One private profile per call: parallel authors would otherwise fight over LibreOffice's single-instance lock.
    with tempfile.TemporaryDirectory(prefix="lo_profile_") as profile:
        cmd = ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--norestore", "--convert-to", fmt,
               "--outdir", str(outdir), str(src)]
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=300, check=False)
    out = outdir / (src.stem + "." + fmt.split(":")[0])
    if r.returncode != 0 or not out.exists():
        raise SourceError(f"LibreOffice failed for {src.name} -> {fmt}: {r.stderr[:300]}")
    return out


def build_scanned_pdf(body: str, dest: Path) -> None:
    """Pages of text drawn as slightly skewed, noisy images (forces the OCR path)."""
    import random

    from PIL import Image, ImageDraw, ImageFilter, ImageFont

    rnd = random.Random(7)
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf", 22)
    plain = re.sub(r"[#*|]", "", body)
    words, lines, cur = plain.split(), [], ""
    for wd in words:
        if len(cur) + len(wd) > 78:
            lines.append(cur)
            cur = ""
        cur += wd + " "
    lines.append(cur)
    pages = []
    for start in range(0, len(lines), 38):
        img = Image.new("L", (1240, 1754), 245)
        d = ImageDraw.Draw(img)
        for k, ln in enumerate(lines[start:start + 38]):
            d.text((90 + rnd.randint(-2, 2), 100 + k * 40), ln, font=font, fill=20)
        img = img.rotate(rnd.uniform(-0.8, 0.8), fillcolor=245).filter(ImageFilter.GaussianBlur(0.6))
        pages.append(img.convert("RGB"))
    pages[0].save(dest, save_all=True, append_images=pages[1:], resolution=150)


# ── pptx / xlsx ────────────────────────────────────────────────────────────────

def build_pptx(body: str, dest: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for chunk in re.split(r"^=== slide\s*$", body, flags=re.MULTILINE):
        if not chunk.strip():
            continue
        notes = ""
        if "\nNotes:" in chunk:
            chunk, notes = chunk.split("\nNotes:", 1)
        lines = [ln.rstrip() for ln in chunk.strip().splitlines() if ln.strip()]
        title = next((ln.lstrip("# ").strip() for ln in lines if ln.startswith("#")), "")
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        tf = slide.placeholders[1].text_frame
        tf.clear()
        first = True
        table_rows = [split_row(ln) for ln in lines if ln.startswith("|") and not TABLE_SEP.match(ln)]
        for ln in lines:
            if ln.startswith(("#", "|")):
                continue
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.text = re.sub(r"^\s*[-*]\s+", "", ln)
        if table_rows:
            cols = max(len(r) for r in table_rows)
            shape = slide.shapes.add_table(len(table_rows), cols, Inches(0.5), Inches(4.2), Inches(9), Inches(0.4 * len(table_rows)))
            for i, r in enumerate(table_rows):
                for j in range(cols):
                    shape.table.cell(i, j).text = r[j] if j < len(r) else ""
        if notes.strip():
            slide.notes_slide.notes_text_frame.text = notes.strip()
    prs.save(dest)


def _cell(v: str):
    s = v.strip()
    if s.startswith("="):
        return s
    try:
        return int(s.replace(",", "")) if re.fullmatch(r"-?\d[\d,]*", s) else float(s) if re.fullmatch(r"-?\d+\.\d+", s) else s
    except ValueError:
        return s


def build_xlsx(body: str, dest: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    sheets = re.split(r"^##\s+", body, flags=re.MULTILINE)
    for block in sheets:
        if not block.strip():
            continue
        name, _, rest = block.partition("\n")
        ws = wb.create_sheet(name.strip()[:31] or "Sheet")
        for ln in rest.splitlines():
            if ln.startswith("|") and not TABLE_SEP.match(ln):
                ws.append([_cell(c) for c in split_row(ln)])
    if not wb.sheetnames:
        raise SourceError(f"{dest.name}: no `## Sheet` blocks")
    wb.save(dest)


# ── driver ─────────────────────────────────────────────────────────────────────

def build_one(src: Path, outdir: Path) -> Path:
    meta, body = parse_source(src)
    out = outdir / meta["out"]
    ext = out.suffix.lstrip(".").lower()
    outdir.mkdir(parents=True, exist_ok=True)
    if ext in VERBATIM:
        out.write_text(body.lstrip("\n"), encoding="utf-8")
    elif ext == "pptx":
        build_pptx(body, out)
    elif ext in ("xlsx", "xls"):
        with tempfile.TemporaryDirectory() as tmp:
            x = Path(tmp) / (out.stem + ".xlsx")
            build_xlsx(body, x)
            if ext == "xlsx":
                shutil.copy(x, out)
            else:
                shutil.copy(soffice(x, "xls", Path(tmp)), out)
    elif ext == "pdf" and meta.get("scan", "false").lower() == "true":
        build_scanned_pdf(body, out)
    elif ext in ("docx", "pdf", "doc", "rtf", "odt"):
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            d = work / (out.stem + ".docx")
            build_docx(body, d, work)
            if ext == "docx":
                shutil.copy(d, out)
            else:
                conv = work / "conv"
                conv.mkdir()
                shutil.copy(soffice(d, ext, conv), out)
    else:
        raise SourceError(f"{src.name}: unsupported output extension .{ext}")
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--lane")
    ap.add_argument("--all", action="store_true")
    ap.add_argument("--check", action="store_true", help="validate front matter only")
    a = ap.parse_args()
    lanes = sorted(p.name for p in (ROOT / "src").iterdir() if p.is_dir()) if a.all else [a.lane] if a.lane else []
    if not lanes:
        ap.error("pass --lane NAME or --all")
    bad = 0
    for lane in lanes:
        for src in sorted((ROOT / "src" / lane).glob("*.src")):
            try:
                if a.check:
                    parse_source(src)
                else:
                    out = build_one(src, ROOT / "files" / lane)
                    print(f"built {out.relative_to(ROOT)} ({out.stat().st_size // 1024} KB)")
            except (SourceError, OSError, subprocess.SubprocessError) as e:
                bad += 1
                print(f"ERROR {src.name}: {e}", file=sys.stderr)
    return 1 if bad else 0


if __name__ == "__main__":
    raise SystemExit(main())
