"""Pictures inside documents: scanned PDF pages, page images, images in Word/PowerPoint files, standalone images.

Text parsers saw only text, so a table pasted as a screenshot, a scanned page inside an otherwise digital PDF or a chart
exported as a picture reached mining as nothing. ``extract`` lists a file's figures in a fixed order (the parser and the
vision reader call it on the same bytes and get the same numbering); ``block`` is the text a parser puts where the figure
sits:

    [Figure 2 · page 3 · picture · read by OCR, not checked]
    ...the OCR text...
    [End of figure 2]

The kind is ``scanned page`` (a page with no text layer: OCR reads it in a usable order) or ``picture`` (a chart, a table
image, a photo). OCR keeps a picture's numbers but not which label each belongs to: mined anyway, a third of the annual
report's chart pairs paired a value with the wrong year (29 of 83). So the mining job reads pictures with the Vision reader
when one is configured (``data/mining/figures.py``) and otherwise holds their OCR text back from mining. Native Office charts are not figures:
their data is read exactly by ``parsers/ooxml_charts.py``.
"""
from __future__ import annotations

import io
import logging
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path

log = logging.getLogger(__name__)

MIN_SIDE_PX = 120          # smaller images are icons, bullets and logos
MIN_BYTES = 4 * 1024
SCANNED_PAGE_CHARS = 40    # a PDF page with less extractable text than this is a picture of a page
PAGE_DPI = 200
SCANNED, PICTURE = "scanned page", "picture"
OCR_MARK = "OCR, not checked"
RENDER_MARK = "rendered text, not checked"
UNCHECKED = frozenset({OCR_MARK, RENDER_MARK})   # block bodies no model has read yet
RENDERED = "rendered page"                       # ``where`` prefix of a page of a file read through LibreOffice
PROSE_CHARS = 200          # a rendered page with at least this much running text is a text page, not a picture
PROSE_WORDS_PER_LINE = 6


def max_render_pages() -> int:
    import os
    raw = os.environ.get("FTS_RENDER_MAX_PAGES", "").strip()
    return int(raw) if raw.isdigit() and int(raw) > 0 else 200

_BLIP_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
_BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
BLOCK_RE = re.compile(r"\[Figure (\d+) · ([^\]]*)\]\n(.*?)\n\[End of figure \1\]", re.DOTALL)


@dataclass
class Figure:
    n: int                 # 1-based, in document order
    where: str             # "page 3", "slide 2", "paragraph 14", "image"
    anchor: int            # page / slide number, or body-child index (docx); 0 for a standalone image
    data: bytes            # PNG or JPEG bytes
    whole_page: bool = False  # a scanned page: the figure IS the page text
    kind_override: str = ""   # a rendered page decides its own kind from its text (``_rendered_kind``)

    @property
    def kind(self) -> str:
        if self.kind_override:
            return self.kind_override
        return SCANNED if self.whole_page else PICTURE

    @property
    def suffix(self) -> str:
        return ".png" if self.data[:8] == b"\x89PNG\r\n\x1a\n" else ".jpg"


def _big_enough(data: bytes) -> bool:
    if len(data) < MIN_BYTES:
        return False
    try:
        from PIL import Image
        with Image.open(io.BytesIO(data)) as im:
            return min(im.size) >= MIN_SIDE_PX
    except Exception:  # noqa: BLE001 - an unreadable image (EMF/WMF, broken) is not a figure we can read
        return False


def _to_png(data: bytes) -> bytes:
    """Formats a vision model or tesseract may not take (TIFF, BMP, GIF, WEBP) as PNG; PNG/JPEG unchanged."""
    if data[:8] == b"\x89PNG\r\n\x1a\n" or data[:3] == b"\xff\xd8\xff":
        return data
    from PIL import Image
    with Image.open(io.BytesIO(data)) as im:
        out = io.BytesIO()
        im.convert("RGB").save(out, "PNG")
        return out.getvalue()


def _docx(path: Path) -> list[Figure]:
    from docx import Document

    doc = Document(str(path))
    out: list[Figure] = []
    for i, child in enumerate(doc.element.body.iterchildren()):
        for blip in child.iter(_BLIP):
            part = doc.part.related_parts.get(blip.get(_BLIP_EMBED, ""))
            data = getattr(part, "blob", b"") if part is not None else b""
            if data and _big_enough(data):
                out.append(Figure(len(out) + 1, f"paragraph {i + 1}", i, _to_png(data)))
    return out


def _pptx(path: Path) -> list[Figure]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def pictures(shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from pictures(sh.shapes)
            elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                yield sh.image.blob

    out: list[Figure] = []
    for s, slide in enumerate(Presentation(str(path)).slides, 1):
        for data in pictures(slide.shapes):
            if _big_enough(data):
                out.append(Figure(len(out) + 1, f"slide {s}", s, _to_png(data)))
    return out


def _render_page(path: Path, page: int) -> bytes:
    pdftoppm = shutil.which("pdftoppm")
    if not pdftoppm:
        raise RuntimeError("pdftoppm (poppler) not found on PATH; install poppler-utils")
    with tempfile.TemporaryDirectory(prefix="_fig_") as tmp:
        base = Path(tmp) / "p"
        r = subprocess.run([pdftoppm, "-r", str(PAGE_DPI), "-png", "-singlefile", "-f", str(page), "-l", str(page),
                            str(path), str(base)], capture_output=True, text=True, timeout=180, check=False)
        png = base.with_suffix(".png")
        if r.returncode != 0 or not png.is_file():
            raise RuntimeError(f"pdftoppm failed on page {page}: {r.stderr.strip()[:200]}")
        return png.read_bytes()


_DO = re.compile(rb"/([^\s/\[\]<>()]+)\s+Do\b")
DECORATION_PAGES = 3       # the same image drawn on this many pages is a logo or letterhead, not content


def _drawn(page) -> set[str] | None:
    """Names of the XObjects the page really draws (``/Im10 Do``), including inside form XObjects; ``None`` if unreadable.

    A resource dictionary is often shared by every page: ``page.images`` then lists every image of the document on
    every page (the annual report's two charts showed up 20 times each)."""
    try:
        seen: set[str] = set()
        contents = page.get_contents()  # a ContentStream is falsy (an empty dict): test for None
        todo = [(contents.get_data() if contents is not None else b"", page.get("/Resources"))]
        while todo:
            data, res = todo.pop()
            xo = res.get_object().get("/XObject") if res is not None else None
            xo = xo.get_object() if xo is not None else {}
            for raw in _DO.findall(data):
                name = raw.decode("latin-1")
                if name in seen:
                    continue
                seen.add(name)
                obj = xo.get(f"/{name}")
                obj = obj.get_object() if obj is not None else None
                if obj is not None and obj.get("/Subtype") == "/Form":
                    todo.append((obj.get_data(), obj.get("/Resources")))
        return seen
    except Exception as e:  # noqa: BLE001 - odd content streams: fall back to every listed image
        log.warning("page content unreadable: %s", e)
        return None


def _pdf(path: Path, page_texts: list[str] | None = None) -> list[Figure]:
    """A page without text is rendered whole (a scan); a page with text contributes the images it draws."""
    import hashlib

    from pypdf import PdfReader

    reader = PdfReader(str(path))
    per_page: list[tuple[int, bool, list[bytes]]] = []
    pages_of: dict[str, set[int]] = {}
    for i, page in enumerate(reader.pages, 1):
        text = page_texts[i - 1] if page_texts and i <= len(page_texts) else (page.extract_text() or "")
        if len(text.strip()) < SCANNED_PAGE_CHARS:
            per_page.append((i, True, [_render_page(path, i)]))
            continue
        drawn = _drawn(page)
        try:
            images = [im for im in page.images if drawn is None or im.name.rsplit(".", 1)[0] in drawn]
        except Exception as e:  # noqa: BLE001 - pypdf raises many types on odd image streams; the text still stands
            log.warning("%s page %d: images unreadable: %s", path.name, i, e)
            images = []
        kept = []
        for im in images:
            try:
                data = im.data
                if _big_enough(data):
                    kept.append(_to_png(data))
            except Exception as e:  # noqa: BLE001 - JBIG2/CCITT streams PIL cannot decode
                log.warning("%s page %d: image %s not decodable: %s", path.name, i, im.name, e)
        for data in kept:
            pages_of.setdefault(hashlib.sha1(data).hexdigest(), set()).add(i)
        per_page.append((i, False, kept))
    out: list[Figure] = []
    for i, whole, datas in per_page:
        for data in datas:
            if not whole and len(pages_of[hashlib.sha1(data).hexdigest()]) >= DECORATION_PAGES:
                continue
            out.append(Figure(len(out) + 1, f"page {i}", i, data, whole_page=whole))
    return out


def _image(path: Path) -> list[Figure]:
    """The image itself (a multi-page TIFF: one scanned page per frame); an image PIL cannot decode is no figure (the
    parse then fails with a reason, not an empty block)."""
    from PIL import Image, ImageSequence

    data = path.read_bytes()
    try:
        with Image.open(io.BytesIO(data)) as im:
            frames = getattr(im, "n_frames", 1)
            if frames > 1 and im.format == "TIFF":
                out = []
                for i, frame in enumerate(ImageSequence.Iterator(im), 1):
                    buf = io.BytesIO()
                    frame.convert("RGB").save(buf, "PNG")
                    out.append(Figure(i, f"page {i}", i, buf.getvalue(), whole_page=True))
                return out
            im.load()
    except Exception as e:  # PIL raises several types for a broken file
        raise ValueError(f"not a readable image: {e}") from e
    return [Figure(1, "image", 0, _to_png(data))]


def _rendered_kind(text: str) -> str:
    """A rendered page with running prose is read like a scanned page; anything else (diagram, chart, form) is a picture."""
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if len(text.strip()) < PROSE_CHARS or not lines:
        return PICTURE
    words = sum(len(ln.split()) for ln in lines) / len(lines)
    return SCANNED if words >= PROSE_WORDS_PER_LINE else PICTURE


@dataclass
class Rendered:
    figures: list[Figure]
    texts: list[str]           # each page's text layer ("" when LibreOffice wrote none)
    pages: int                 # pages of the rendered PDF (more than ``figures`` when capped)


def rendered(path: Path) -> Rendered:
    """Every page of ``path`` rendered by LibreOffice, as figures (capped at ``max_render_pages``). Raises ``RenderError``."""
    from pypdf import PdfReader

    from finetune_studio.data.render import convert

    with tempfile.TemporaryDirectory(prefix="_render_") as tmp:
        src = Path(tmp) / f"doc{path.suffix.lower()}"     # a name LibreOffice cannot misread (spaces, unicode)
        shutil.copyfile(path, src)
        pdf = convert(src, "pdf", Path(tmp))
        reader = PdfReader(str(pdf))
        n = len(reader.pages)
        figs_out, texts = [], []
        for i, page in enumerate(reader.pages[:max_render_pages()], 1):
            text = (page.extract_text() or "").strip()
            texts.append(text)
            figs_out.append(Figure(i, f"{RENDERED} {i}", i, _render_page(pdf, i), whole_page=True,
                                   kind_override=_rendered_kind(text)))
        return Rendered(figs_out, texts, n)


_BY_SUFFIX = {".docx": _docx, ".pptx": _pptx, ".pdf": _pdf,
              ".png": _image, ".jpg": _image, ".jpeg": _image, ".tif": _image, ".tiff": _image, ".bmp": _image,
              ".webp": _image, ".gif": _image}


def extract(path: Path, page_texts: list[str] | None = None) -> list[Figure]:
    """The file's figures in document order; ``[]`` for a format without pictures (or a file we cannot open).

    ``page_texts``: a PDF's per-page text as the parser extracted it, so the parser and this agree on scanned pages."""
    fn = _BY_SUFFIX.get(path.suffix.lower())
    if fn is None:
        return []
    try:
        return fn(path, page_texts) if fn is _pdf else fn(path)
    except (zipfile.BadZipFile, OSError, ValueError, KeyError) as e:
        log.warning("figures of %s unreadable: %s", path.name, e)
        return []


def ocr(fig: Figure, languages: str) -> str:
    """Tesseract text of a figure ("" when it found none)."""
    from finetune_studio.data.ocr import ocr_image

    with tempfile.NamedTemporaryFile(suffix=fig.suffix) as f:
        f.write(fig.data)
        f.flush()
        return ocr_image(f.name, languages=languages).strip()


def block(fig: Figure, body: str, read_by: str) -> str:
    """The text a parser or the vision reader puts where the figure sits."""
    body = body.strip() or "(no text found)"
    return f"[Figure {fig.n} · {fig.where} · {fig.kind} · read by {read_by}]\n{body}\n[End of figure {fig.n}]"


def header_parts(header: str) -> tuple[str, str, str]:
    """(where, kind, read_by) of a block header ``page 3 · picture · read by OCR, not checked``."""
    parts = header.split(" · ")
    where = parts[0]
    kind = parts[1] if len(parts) >= 3 else PICTURE
    read_by = parts[-1].removeprefix("read by ")
    return where, kind, read_by


def rendered_block(fig: Figure, text_layer: str, languages: str, warnings: list[str]) -> str:
    """A rendered page: its PDF text layer when LibreOffice wrote one, else OCR; unchecked either way."""
    if text_layer.strip():
        return block(fig, text_layer, RENDER_MARK)
    return ocr_block(fig, languages, warnings)


def ocr_block(fig: Figure, languages: str, warnings: list[str]) -> str:
    try:
        text = ocr(fig, languages)
    except Exception as e:  # noqa: BLE001 - the figure is still listed; mining's vision reader may read it
        warnings.append(f"figure {fig.n} ({fig.where}): OCR failed: {e}")
        text = ""
    return block(fig, text, OCR_MARK)


def summary(figs: list[Figure]) -> list[dict]:
    """``structured["figures"]`` of a parse: where each figure is, never its bytes."""
    return [{"n": f.n, "where": f.where, "whole_page": f.whole_page, "bytes": len(f.data)} for f in figs]
