"""XLSX (modern Excel) parser."""

from __future__ import annotations

from pathlib import Path

from ._base import cell_text, cli_run, make_result
from .ooxml_charts import charts_in_package, sheet_resolver


def parse(path: Path) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        return make_result("", {"type": "xlsx", "error": str(e)}, parser="xlsx_v1",
                           warnings=["install openpyxl (pip install openpyxl)"])
    wb = load_workbook(str(path), data_only=True, read_only=True)
    try:
        text_lines = []
        sheets = []
        grid: dict[str, list[list[str]]] = {}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            grid[sheet_name] = []
            for row in ws.iter_rows(min_row=1, min_col=1, values_only=True):
                values = [cell_text(v) for v in row]
                grid[sheet_name].append(values)  # by position, for chart references
                if any(v.strip() for v in values):
                    rows.append(values)
            sheets.append({"name": sheet_name, "row_count": len(rows), "rows": rows[:200]})
            text_lines.append(f"=== Sheet: {sheet_name} ({len(rows)} rows) ===")
            for row in rows:
                text_lines.append(" | ".join(row))
        charts = charts_in_package(path, sheet_resolver(grid))  # openpyxl does not read charts back; their cached data is in the package
        if charts:
            text_lines.append("=== Charts ===")
            text_lines.extend(charts)
        text = "\n".join(text_lines)
        structured = {
            "chart_count": len(charts),
            "type": "xlsx",
            "sheet_count": len(sheets),
            "sheet_names": wb.sheetnames,
            "sheets": sheets,
        }
    finally:
        wb.close()
    return make_result(text, structured, parser="xlsx_v1")


if __name__ == "__main__":
    cli_run(parse)
