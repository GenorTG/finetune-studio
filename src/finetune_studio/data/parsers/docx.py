"""DOCX (modern Word) parser."""

from __future__ import annotations

from pathlib import Path

from ._base import cli_run, make_result


def parse(path: Path) -> dict:
    try:
        from docx import Document
    except ImportError as e:
        return make_result("", {"type": "docx", "error": str(e)}, parser="docx_v1",
                           warnings=["install python-docx (pip install python-docx)"])
    doc = Document(str(path))
    paragraphs = []
    headings = []
    for p in doc.paragraphs:
        if not p.text.strip():
            continue
        paragraphs.append(p.text)
        if p.style and p.style.name and p.style.name.startswith("Heading"):
            try:
                level = int(p.style.name.split()[-1])
            except ValueError:
                level = 0
            headings.append({"level": level, "text": p.text.strip()})
    tables = []
    for t in doc.tables:
        rows = []
        for row in t.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append({"rows": rows})
    text = "\n".join(paragraphs)
    structured = {
        "type": "docx",
        "paragraph_count": len(paragraphs),
        "headings": headings,
        "table_count": len(tables),
        "tables": tables[:50],
    }
    return make_result(text, structured, parser="docx_v1")


if __name__ == "__main__":
    cli_run(parse)
