"""PPTX (PowerPoint) parser."""

from __future__ import annotations

from pathlib import Path

from .. import figures as figs
from ..ocr import DEFAULT_LANGS
from ._base import cell_text, cli_run, make_result
from .ooxml_charts import chart_table
from .ooxml_diagrams import REL_ATTR, REL_IDS_TAG, diagram_outline


def parse(path: Path) -> dict:
    try:
        from pptx import Presentation
    except ImportError as e:
        return make_result("", {"type": "pptx", "error": str(e)}, parser="pptx_v1",
                           warnings=["install python-pptx (pip install python-pptx)"])
    prs = Presentation(str(path))
    warnings: list[str] = []
    found = figs.extract(path)
    text_lines = []
    slides_structured = []
    n_charts = 0
    n_diagrams = 0
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text)
            if getattr(shape, "has_chart", False) and (chart := chart_table(shape.chart.part.blob)):
                texts.append(chart)
                n_charts += 1
            for ref in shape._element.iter(REL_IDS_TAG):  # SmartArt graphic frame: outline from its data part
                try:
                    blob = slide.part.related_part(ref.get(REL_ATTR, "")).blob
                except KeyError:
                    continue
                if outline := diagram_outline(blob):
                    texts.append(outline)
                    n_diagrams += 1
            if shape.has_table:
                for row in shape.table.rows:
                    texts.append(" | ".join(cell_text(cell.text) for cell in row.cells))
        notes_text = ""
        if slide.has_notes_slide:
            notes_shape = slide.notes_slide.notes_text_frame
            notes_text = notes_shape.text.strip() if notes_shape is not None else ""
            if notes_text:
                texts.append(f"[Speaker notes] {notes_text}")
        texts.extend(figs.ocr_block(f, DEFAULT_LANGS, warnings) for f in found if f.anchor == i)
        slide_text = "\n".join(texts)
        text_lines.append(f"=== Slide {i} ===\n{slide_text}")
        slides_structured.append({"slide": i, "shape_count": len(slide.shapes), "text": slide_text,
                                  "notes": notes_text})
    text = "\n\n".join(text_lines)
    structured = {"type": "pptx", "slide_count": len(prs.slides), "slides": slides_structured,
                  "chart_count": n_charts, "diagram_count": n_diagrams, "figures": figs.summary(found)}
    return make_result(text, structured, parser="pptx_v1", warnings=warnings)


if __name__ == "__main__":
    cli_run(parse)
