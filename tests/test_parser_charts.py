"""Native Office charts reach the parsed text as data tables (``data/parsers/ooxml_charts.py``)."""
from __future__ import annotations

from pathlib import Path

import pytest


def test_pptx_chart_values_are_in_the_slide_text(tmp_path: Path) -> None:
    pptx = pytest.importorskip("pptx")
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    from finetune_studio.data.parsers import pptx as parser

    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Revenue"
    data = CategoryChartData()
    data.categories = ["Q1", "Q2", "Q3"]
    data.add_series("2023", (24.1, 25.3, 27.9))
    data.add_series("2024", (26.0, 28.4, 30.2))
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(6), Inches(4), data)
    gf.chart.has_title = True
    gf.chart.chart_title.text_frame.text = "Revenue by quarter"
    path = tmp_path / "deck.pptx"
    prs.save(str(path))

    out = parser.parse(path)
    text = out["text"]
    assert "=== Slide 1 ===" in text and "[Chart: Revenue by quarter (bar chart)]" in text
    assert "Category | 2023 | 2024" in text and "Q2 | 25.3 | 28.4" in text
    assert out["structured"]["chart_count"] == 1


def test_xlsx_chart_is_appended_with_its_cached_values(tmp_path: Path) -> None:
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl.chart import LineChart, Reference

    from finetune_studio.data.parsers import xlsx as parser

    wb = openpyxl.Workbook()
    ws = wb.active
    for row in [("Month", "Excursions"), ("Jan", 4), ("Feb", 7), ("Mar", 2)]:
        ws.append(row)
    ch = LineChart()
    ch.title = "Excursions per month"
    ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4))
    ws.add_chart(ch, "D2")
    path = tmp_path / "book.xlsx"
    wb.save(str(path))

    out = parser.parse(path)
    # openpyxl writes only the formula references (no value cache): the values are read from the cells
    assert out["structured"]["chart_count"] == 1
    assert "[Chart: Excursions per month (line chart)]\nCategory | Excursions\nJan | 4\nFeb | 7\nMar | 2" in out["text"]


def test_chart_table_reads_cached_points_in_index_order() -> None:
    from finetune_studio.data.parsers.ooxml_charts import chart_table

    xml = b"""<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"
      xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><c:chart>
      <c:title><c:tx><c:rich><a:p><a:r><a:t>Net debt</a:t></a:r></a:p></c:rich></c:tx></c:title>
      <c:plotArea><c:pieChart><c:ser><c:tx><c:strRef><c:strCache><c:pt idx="0"><c:v>EUR m</c:v></c:pt></c:strCache></c:strRef></c:tx>
      <c:cat><c:strRef><c:strCache><c:pt idx="1"><c:v>Lease</c:v></c:pt><c:pt idx="0"><c:v>Term loan</c:v></c:pt></c:strCache></c:strRef></c:cat>
      <c:val><c:numRef><c:numCache><c:pt idx="0"><c:v>38.5</c:v></c:pt><c:pt idx="1"><c:v>21.7</c:v></c:pt></c:numCache></c:numRef></c:val>
      </c:ser></c:pieChart></c:plotArea></c:chart></c:chartSpace>"""
    assert chart_table(xml) == "[Chart: Net debt (pie chart)]\nCategory | EUR m\nTerm loan | 38.5\nLease | 21.7"
    assert chart_table(b"not xml") == ""


def test_docx_inline_chart_appears_where_it_sits(tmp_path: Path) -> None:
    docx = pytest.importorskip("docx")
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from docx.oxml import parse_xml

    from finetune_studio.data.parsers import docx as parser

    chart_xml = (b'<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"><c:chart><c:plotArea>'
                 b'<c:barChart><c:ser><c:tx><c:v>Pallets</c:v></c:tx><c:cat><c:strLit><c:pt idx="0"><c:v>40ft</c:v></c:pt>'
                 b'<c:pt idx="1"><c:v>20ft</c:v></c:pt></c:strLit></c:cat><c:val><c:numLit><c:pt idx="0"><c:v>23</c:v></c:pt>'
                 b'<c:pt idx="1"><c:v>10</c:v></c:pt></c:numLit></c:val></c:ser></c:barChart></c:plotArea></c:chart></c:chartSpace>')
    d = docx.Document()
    d.add_paragraph("Loading limits:")
    part = Part(PackURI("/word/charts/chart1.xml"),
                "application/vnd.openxmlformats-officedocument.drawingml.chart+xml", chart_xml, d.part.package)
    rid = d.part.relate_to(part, RT.CHART)
    p = d.add_paragraph()
    p._p.append(parse_xml(
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:drawing>'
        '<c:chart xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="{rid}"/>'
        '</w:drawing></w:r>'))
    d.add_paragraph("Pharma loads: 20.")
    path = tmp_path / "sop.docx"
    d.save(str(path))

    out = parser.parse(path)
    assert out["structured"]["chart_count"] == 1
    assert "Loading limits:\n[Chart: untitled (bar chart)]\nCategory | Pallets\n40ft | 23\n20ft | 10\nPharma loads: 20." in out["text"]
