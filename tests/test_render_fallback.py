"""Unsupported or text-less files are opened with LibreOffice and read page by page (``data/render.py``); SmartArt and
legacy .xls charts are read exactly (``parsers/ooxml_diagrams.py``, ``parsers/xls.py``)."""
from __future__ import annotations

import io
import shutil
from pathlib import Path

import pytest

from finetune_studio.data import figures
from finetune_studio.data.parsers import parse

needs_lo = pytest.mark.skipif(not (shutil.which("soffice") and shutil.which("pdftoppm")),
                              reason="LibreOffice/poppler not installed")

SVG = ('<svg xmlns="http://www.w3.org/2000/svg" width="800" height="400"><rect x="20" y="40" width="300" height="100" '
       'fill="none" stroke="black"/><text x="40" y="100" font-size="28">Pump P-12</text><text x="480" y="100" '
       'font-size="28">Tank T-3</text><text x="300" y="200" font-size="24">feeds at 40 l/min</text></svg>')


@needs_lo
def test_a_file_without_a_parser_is_rendered_and_read_page_by_page(tmp_path: Path) -> None:
    from finetune_studio.data.render import convert
    svg = tmp_path / "plant.svg"
    svg.write_text(SVG)
    odg = convert(svg, "odg", tmp_path)                 # .odg has no parser
    out = parse(odg)
    assert out["metadata"]["parser"] == "render_v1_libreoffice"
    assert "no parser for .odg" in out["metadata"]["warnings"][0]
    assert "[Figure 1 · rendered page 1 · picture · read by rendered text, not checked]" in out["text"]
    assert "Pump P-12" in out["text"] and "40 l/min" in out["text"]
    assert out["structured"]["figures"][0]["where"] == "rendered page 1"


@needs_lo
def test_thin_parser_text_falls_back_to_rendering(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    docx = pytest.importorskip("docx")
    from finetune_studio.data import parsers
    from finetune_studio.data.parsers._base import make_result
    d = docx.Document()
    for _ in range(8):
        d.add_paragraph("The reefer pre-cool must reach the set-point before any pallet is loaded into the box.")
    p = tmp_path / "sop.docx"
    d.save(str(p))
    monkeypatch.setattr(parsers, "get_parser_for", lambda path: (lambda q: make_result("", {}, parser="docx_v1")))
    out = parse(p)
    assert out["metadata"]["parser"] == "render_v1_libreoffice"
    assert "almost no text" in out["metadata"]["warnings"][0]
    assert "· scanned page · read by rendered text" in out["text"]       # running prose: mined like a scanned page
    assert "pre-cool must reach the set-point" in out["text"]


@needs_lo
def test_a_crashing_parser_falls_back_to_rendering(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from finetune_studio.data import parsers
    from finetune_studio.data.render import convert
    svg = tmp_path / "plant.svg"
    svg.write_text(SVG)
    odt = convert(svg, "odg", tmp_path)

    def boom(path: Path) -> dict:
        raise ValueError("corrupt stream")

    monkeypatch.setattr(parsers, "get_parser_for", lambda path: boom)
    out = parse(odt)
    assert out["metadata"]["parser"] == "render_v1_libreoffice" and "parser failed (ValueError: corrupt stream)" in \
        out["metadata"]["warnings"][0]


def test_binary_without_parser_is_never_read_as_text(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from finetune_studio.data import render
    p = tmp_path / "blob.xyz"
    p.write_bytes(b"\x00\x01\x02garbage" * 100)
    monkeypatch.setattr(render, "soffice", lambda: None)
    out = parse(p)
    assert out["text"] == "" and "LibreOffice not found" in out["metadata"]["warnings"][0]
    t = tmp_path / "notes.weird"
    t.write_text("plain notes with an odd extension")
    assert parse(t)["metadata"]["parser"] == "text_fallback"


def test_multipage_tiff_is_one_scanned_page_per_frame(tmp_path: Path) -> None:
    from PIL import Image
    frames = [Image.new("RGB", (600, 400), c) for c in ("white", "white", "white")]
    p = tmp_path / "scan.tiff"
    frames[0].save(p, save_all=True, append_images=frames[1:])
    figs = figures.extract(p)
    assert [(f.where, f.kind) for f in figs] == [("page 1", "scanned page"), ("page 2", "scanned page"),
                                                 ("page 3", "scanned page")]


def test_rendered_page_kind() -> None:
    prose = "\n".join(["The reefer pre-cool must reach the set-point before any pallet is loaded."] * 4)
    assert figures._rendered_kind(prose) == figures.SCANNED
    assert figures._rendered_kind("Pump P-12 Tank T-3\nfeeds at 40 l/min") == figures.PICTURE
    labels = "\n".join(["Q1", "24.1", "Q2", "25.3"] * 20)
    assert figures._rendered_kind(labels) == figures.PICTURE


SMARTART = b"""<dgm:dataModel xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram"
 xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><dgm:ptLst>
 <dgm:pt modelId="0" type="doc"/>
 <dgm:pt modelId="1"><dgm:t><a:p><a:r><a:t>Duty Ops</a:t></a:r></a:p></dgm:t></dgm:pt>
 <dgm:pt modelId="2"><dgm:t><a:p><a:r><a:t>Fleet</a:t></a:r></a:p></dgm:t></dgm:pt>
 <dgm:pt modelId="3"><dgm:t><a:p><a:r><a:t>Quality on-call</a:t></a:r></a:p></dgm:t></dgm:pt>
 <dgm:pt modelId="9" type="parTrans"><dgm:t><a:p><a:r><a:t>ignored</a:t></a:r></a:p></dgm:t></dgm:pt>
 </dgm:ptLst><dgm:cxnLst>
 <dgm:cxn modelId="c1" srcId="0" destId="1" srcOrd="0"/>
 <dgm:cxn modelId="c2" srcId="1" destId="2" srcOrd="1"/>
 <dgm:cxn modelId="c3" srcId="1" destId="3" srcOrd="0"/>
 <dgm:cxn modelId="c4" type="presOf" srcId="1" destId="9"/>
 </dgm:cxnLst></dgm:dataModel>"""


def test_smartart_outline_follows_the_hierarchy() -> None:
    from finetune_studio.data.parsers.ooxml_diagrams import diagram_outline
    assert diagram_outline(SMARTART) == "[Diagram]\n- Duty Ops\n  - Quality on-call\n  - Fleet"
    assert diagram_outline(b"<broken") == ""


def test_docx_smartart_sits_where_it_is(tmp_path: Path) -> None:
    docx = pytest.importorskip("docx")
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from docx.oxml import parse_xml

    from finetune_studio.data.parsers import docx as parser
    d = docx.Document()
    d.add_paragraph("Escalation:")
    part = Part(PackURI("/word/diagrams/data1.xml"),
                "application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml", SMARTART, d.part.package)
    rid = d.part.relate_to(part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramData")
    d.add_paragraph()._p.append(parse_xml(
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:drawing>'
        '<dgm:relIds xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:dm="{rid}"/>'
        '</w:drawing></w:r>'))
    d.add_paragraph("Call in this order.")
    p = tmp_path / "esc.docx"
    d.save(str(p))
    out = parser.parse(p)
    assert "Escalation:\n[Diagram]\n- Duty Ops\n  - Quality on-call\n  - Fleet\nCall in this order." in out["text"]
    assert out["structured"]["diagram_count"] == 1


@needs_lo
def test_legacy_xls_chart_is_read_through_libreoffice(tmp_path: Path) -> None:
    openpyxl = pytest.importorskip("openpyxl")
    pytest.importorskip("xlrd")
    from openpyxl.chart import BarChart, Reference

    from finetune_studio.data.parsers import xls as parser
    from finetune_studio.data.render import convert
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Alerts"
    for row in [("Half", "Alerts"), ("H1 2023", 5.2), ("H2 2023", 2.7), ("H1 2024", 1.9)]:
        ws.append(row)
    ch = BarChart()
    ch.title = "Excursion alerts"
    ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4))
    ws.add_chart(ch, "D2")
    src = tmp_path / "alerts.xlsx"
    wb.save(str(src))
    xls = convert(src, "xls", tmp_path / "out") if (tmp_path / "out").mkdir() is None else None
    out = parser.parse(xls)
    assert out["structured"]["chart_count"] == 1, out["metadata"]["warnings"]
    assert "Excursion alerts" in out["text"] and "H2 2023 | 2.7" in out["text"]


def test_multipage_tiff_parse_has_a_block_per_page(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from PIL import Image

    from finetune_studio.data.parsers import image
    monkeypatch.setattr(figures, "ocr", lambda fig, languages: f"text of {fig.where}")
    p = tmp_path / "scan.tiff"
    frames = [Image.new("RGB", (600, 400), "white") for _ in range(2)]
    frames[0].save(p, save_all=True, append_images=frames[1:])
    text = image.parse(p)["text"]
    assert "[Figure 1 · page 1 · scanned page" in text and "text of page 2" in text and "[End of figure 2]" in text
