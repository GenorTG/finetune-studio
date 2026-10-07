"""Scanned PDFs are OCR'd without optional Python wrappers, and legacy .doc keeps its tables.

Both were silent data loss found by the corpus smoke test: a scanned PDF parsed to 0 characters (pdf2image/pytesseract
missing) and a table vanished from a .doc (no antiword/catdoc, crude olefile scan).
"""
from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

import pytest

from finetune_studio.data.parsers import doc as doc_parser
from finetune_studio.data.parsers import pdf as pdf_parser

needs_ocr = pytest.mark.skipif(not (shutil.which("tesseract") and shutil.which("pdftoppm")), reason="tesseract/poppler missing")
needs_office = pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice missing")


def _scanned_pdf(path: Path, line: str) -> None:
    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("RGB", (1240, 400), "white")
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 34)
    ImageDraw.Draw(img).text((60, 120), line, font=font, fill="black")
    img.save(path, resolution=150)


@needs_ocr
def test_scanned_pdf_is_ocrd_through_the_cli(tmp_path: Path) -> None:
    pdf = tmp_path / "scan.pdf"
    _scanned_pdf(pdf, "Breaker in bay 7 trips at 41 amps")
    result = pdf_parser.parse(pdf)
    assert "41 amps" in result["text"]
    assert "ocr" in result["metadata"]["parser"]


def test_missing_poppler_is_an_error_not_an_empty_document(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from finetune_studio.data import ocr

    monkeypatch.setattr(ocr.shutil, "which", lambda name: None)
    with pytest.raises(RuntimeError, match="pdftoppm"):
        ocr.ocr_pdf(tmp_path / "x.pdf")


@needs_office
def test_legacy_doc_keeps_its_table(tmp_path: Path) -> None:
    from docx import Document

    d = Document()
    d.add_paragraph("Rate sheet for the Gdynia depot")
    t = d.add_table(rows=3, cols=2)
    for i, (a, b) in enumerate([("item", "price"), ("plug rental", "13.40 EUR"), ("gate fee", "22 EUR")]):
        t.cell(i, 0).text, t.cell(i, 1).text = a, b
    src = tmp_path / "rates.docx"
    d.save(src)
    subprocess.run(["soffice", f"-env:UserInstallation=file://{tmp_path}/p", "--headless", "--convert-to", "doc",
                    "--outdir", str(tmp_path), str(src)], capture_output=True, timeout=180, check=True)
    text = doc_parser.parse(tmp_path / "rates.doc")["text"]
    assert "plug rental" in text and "13.40 EUR" in text and "gate fee" in text
