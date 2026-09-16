"""Durable coverage: every extension in finetune_studio.data.parsers.PARSERS.

Text / structured formats use minimal stdlib fixtures and must extract content.
Binary / OCR formats are dependency-aware: when the Python (or system) deps are
missing we assert clear warnings or skip — never unexpected exceptions.
"""
from __future__ import annotations

import importlib
import importlib.util
import io
import struct
import tomllib
import zipfile
from pathlib import Path

import pytest

from finetune_studio.data.parsers import PARSERS, parse

MARKER = "HeliosMarkUniqueToken42"

# Extensions that need optional Python packages to produce real content.
_PYTHON_DEP_EXTS: frozenset[str] = frozenset({
    ".pdf", ".docx", ".xlsx", ".xls", ".pptx",
    ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp", ".gif",
})
# Legacy DOC / Outlook MSG need system or specialised tooling.
_SYSTEM_DEP_EXTS: frozenset[str] = frozenset({".doc", ".msg"})

_MISSING_DEP_HINTS: tuple[str, ...] = (
    "install ",
    "not installed",
    "OCR failed",
    "no legacy-DOC",
    "No text extracted",
    "pypdf",
    "poppler",
    "tesseract",
    "antiword",
    "beautifulsoup4",
    "striprtf",
)


def _has_module(name: str) -> bool:
    return importlib.util.find_spec(name) is not None


def _write_text(path: Path, body: str) -> None:
    path.write_text(body, encoding="utf-8")


def _minimal_png(path: Path) -> None:
    """1x1 opaque PNG without Pillow (stdlib zlib + CRC)."""
    import zlib

    def chunk(tag: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + tag
            + data
            + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
        )

    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    raw = zlib.compress(b"\x00\xff\x00\x00")
    path.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", raw)
        + chunk(b"IEND", b"")
    )


def _write_odf(path: Path, ext: str) -> None:
    ns = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"
    content = (
        f'<?xml version="1.0"?>'
        f'<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        f'xmlns:text="{ns}">'
        f'<office:body><office:text>'
        f'<text:p xmlns:text="{ns}">{MARKER}</text:p>'
        f"</office:text></office:body></office:document-content>"
    )
    mime = {
        ".odt": "application/vnd.oasis.opendocument.text",
        ".ods": "application/vnd.oasis.opendocument.spreadsheet",
        ".odp": "application/vnd.oasis.opendocument.presentation",
    }[ext]
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("mimetype", mime, compress_type=zipfile.ZIP_STORED)
        zf.writestr("content.xml", content)


def _write_epub(path: Path) -> None:
    container = (
        '<?xml version="1.0"?>'
        '<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
        "<rootfiles><rootfile full-path=\"OEBPS/chap.xhtml\" "
        'media-type="application/xhtml+xml"/></rootfiles></container>'
    )
    xhtml = (
        f'<?xml version="1.0"?><html xmlns="http://www.w3.org/1999/xhtml">'
        f"<body><p>{MARKER}</p></body></html>"
    )
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        zf.writestr("META-INF/container.xml", container)
        zf.writestr("OEBPS/chap.xhtml", xhtml)


def _write_pdf(path: Path) -> None:
    """Minimal PDF with a Helvetica text show operator (extractable by pypdf)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({MARKER}) Tj ET"
    objects = [
        "1 0 obj<< /Type /Catalog /Pages 2 0 R >>endobj\n",
        "2 0 obj<< /Type /Pages /Kids [3 0 R] /Count 1 >>endobj\n",
        (
            "3 0 obj<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            "/Contents 4 0 R /Resources<< /Font<< /F1 5 0 R >> >> >>endobj\n"
        ),
        f"4 0 obj<< /Length {len(stream)} >>stream\n{stream}\nendstream\nendobj\n",
        "5 0 obj<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>endobj\n",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = [0]
    for obj in objects:
        offsets.append(out.tell())
        out.write(obj.encode("latin-1"))
    xref_pos = out.tell()
    out.write(f"xref\n0 {len(offsets)}\n".encode("latin-1"))
    out.write(b"0000000000 65535 f \n")
    for off in offsets[1:]:
        out.write(f"{off:010d} 00000 n \n".encode("latin-1"))
    out.write(
        (
            f"trailer<< /Size {len(offsets)} /Root 1 0 R >>\n"
            f"startxref\n{xref_pos}\n%%EOF\n"
        ).encode("latin-1")
    )
    path.write_bytes(out.getvalue())


def _write_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph(MARKER)
    doc.save(str(path))


def _write_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    wb.active["A1"] = MARKER
    wb.save(str(path))


def _write_pptx(path: Path) -> None:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = MARKER
    prs.save(str(path))


def _write_xls(path: Path) -> None:
    # xlrd is read-only; use xlwt when present, else skip creation.
    xlwt = pytest.importorskip("xlwt", reason="xlwt needed to build .xls fixture")
    book = xlwt.Workbook()
    sheet = book.add_sheet("s1")
    sheet.write(0, 0, MARKER)
    book.save(str(path))


def _fixture_for(ext: str, dest: Path) -> str:
    """Write a minimal fixture for ``ext``. Returns mode: expect_text|dep_warn|skip."""
    dest.parent.mkdir(parents=True, exist_ok=True)
    text_body = f"Document about Helios. Marker={MARKER}\n"

    if ext in {".txt", ".md", ".markdown", ".log", ".py", ".js", ".ts", ".jsx", ".tsx",
               ".css", ".yaml", ".yml", ".ini", ".cfg", ".conf"}:
        _write_text(dest, text_body)
        return "expect_text"

    if ext in {".csv", ".tsv"}:
        sep = "\t" if ext == ".tsv" else ","
        _write_text(dest, f"key{sep}value\nmarker{sep}{MARKER}\n")
        return "expect_text"

    if ext == ".json":
        _write_text(dest, f'{{"marker": "{MARKER}"}}\n')
        return "expect_text"

    if ext == ".jsonl":
        _write_text(dest, f'{{"marker": "{MARKER}"}}\n{{"n": 1}}\n')
        return "expect_text"

    if ext == ".xml":
        _write_text(dest, f"<root><item>{MARKER}</item></root>\n")
        return "expect_text"

    if ext in {".html", ".htm"}:
        _write_text(dest, f"<html><body><p>{MARKER}</p></body></html>\n")
        return "expect_text"

    if ext == ".rtf":
        _write_text(dest, r"{\rtf1\ansi\deff0 " + MARKER + r"\par}")
        return "expect_text"

    if ext == ".eml":
        _write_text(
            dest,
            (
                "From: a@example.com\nTo: b@example.com\nSubject: marker\n"
                f"MIME-Version: 1.0\nContent-Type: text/plain; charset=utf-8\n\n{MARKER}\n"
            ),
        )
        return "expect_text"

    if ext in {".odt", ".ods", ".odp"}:
        _write_odf(dest, ext)
        return "expect_text"

    if ext == ".epub":
        _write_epub(dest)
        return "expect_text"

    if ext == ".pdf":
        if not _has_module("pypdf"):
            dest.write_bytes(b"%PDF-1.1\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF\n")
            return "dep_warn"
        _write_pdf(dest)
        return "expect_text"

    if ext == ".docx":
        if not _has_module("docx"):
            dest.write_bytes(b"PK\x03\x04not-a-real-docx")
            return "dep_warn"
        _write_docx(dest)
        return "expect_text"

    if ext == ".xlsx":
        if not _has_module("openpyxl"):
            dest.write_bytes(b"PK\x03\x04not-a-real-xlsx")
            return "dep_warn"
        _write_xlsx(dest)
        return "expect_text"

    if ext == ".xls":
        if not _has_module("xlrd"):
            dest.write_bytes(b"not-a-real-xls")
            return "dep_warn"
        if not _has_module("xlwt"):
            pytest.skip("xlrd installed but xlwt unavailable to build .xls fixture")
        _write_xls(dest)
        return "expect_text"

    if ext == ".pptx":
        if not _has_module("pptx"):
            dest.write_bytes(b"PK\x03\x04not-a-real-pptx")
            return "dep_warn"
        _write_pptx(dest)
        return "expect_text"

    if ext in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp", ".gif"}:
        # 1x1 PNG bytes under any image extension — OCR on a blank pixel is
        # not reliable, so we always treat this as a dep_warn / no-crash path.
        _minimal_png(dest)
        return "dep_warn"

    if ext == ".doc":
        dest.write_bytes(b"\xd0\xcf\x11\xe0" + b"\x00" * 64)  # OLE-ish stub
        return "dep_warn"

    if ext == ".msg":
        dest.write_bytes(b"\xd0\xcf\x11\xe0" + b"\x00" * 64)
        return "dep_warn"

    pytest.fail(f"no fixture builder for extension {ext}")


def _warnings_look_clear(warnings: list[str]) -> bool:
    blob = " ".join(warnings).lower()
    return any(h.lower() in blob for h in _MISSING_DEP_HINTS)


class TestParsersExtraDeclared:
    def test_pyproject_parsers_extra_and_all(self) -> None:
        root = Path(__file__).resolve().parents[1]
        data = tomllib.loads((root / "pyproject.toml").read_text(encoding="utf-8"))
        extras = data["project"]["optional-dependencies"]
        parsers = extras["parsers"]
        expected = {
            "pypdf", "python-docx", "openpyxl", "xlrd",
            "python-pptx", "beautifulsoup4", "striprtf", "Pillow",
        }
        names = {p.split(">=")[0].split("[")[0] for p in parsers}
        assert names == expected
        # Must not pretend system tools are pip packages.
        joined = " ".join(parsers).lower()
        for forbidden in ("antiword", "tesseract", "poppler", "pdftotext"):
            assert forbidden not in joined
        all_extra = extras["all"]
        for pkg in expected:
            assert any(pkg in item for item in all_extra), f"{pkg} missing from all"


@pytest.mark.parametrize("ext", sorted(PARSERS.keys()))
def test_every_registered_extension(ext: str, tmp_path: Path) -> None:
    assert ext in PARSERS
    path = tmp_path / f"sample{ext}"
    mode = _fixture_for(ext, path)
    assert path.exists()

    result = parse(path)

    assert isinstance(result, dict)
    assert "text" in result
    assert "metadata" in result
    warnings = list(result.get("metadata", {}).get("warnings") or [])

    if mode == "expect_text":
        assert MARKER in (result.get("text") or ""), (
            f"{ext}: expected marker in text; warnings={warnings!r}"
        )
        return

    # Binary / OCR / missing-dep paths: no crash; content, clear warning, or
    # best-effort non-empty extraction (e.g. .msg OLE read as text) is OK.
    text = (result.get("text") or "").strip()
    if MARKER in text:
        return
    if ext in _SYSTEM_DEP_EXTS | _PYTHON_DEP_EXTS or mode == "dep_warn":
        if text:
            return
        assert warnings, f"{ext}: empty text without warnings"
        assert _warnings_look_clear(warnings), (
            f"{ext}: warnings not clearly dependency-related: {warnings!r}"
        )
        return
    pytest.fail(f"{ext}: unexpected empty parse; warnings={warnings!r}")


def test_html_missing_bs4_warns_or_extracts(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "x.html"
    path.write_text(f"<html><body><p>{MARKER}</p></body></html>", encoding="utf-8")
    if _has_module("bs4"):
        result = parse(path)
        assert MARKER in result["text"]
        return
    result = parse(path)
    assert MARKER in result["text"]
    assert any("beautifulsoup4" in w.lower() for w in result["metadata"]["warnings"])
