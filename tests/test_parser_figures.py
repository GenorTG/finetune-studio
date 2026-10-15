"""Pictures in documents become figure blocks with OCR text (``data/figures.py``).

Needs the tesseract and pdftoppm binaries (skipped without them), like the real parse path.
"""
from __future__ import annotations

import io
import shutil
from pathlib import Path

import pytest

from finetune_studio.data import figures

pytestmark = pytest.mark.skipif(not (shutil.which("tesseract") and shutil.which("pdftoppm")),
                                reason="tesseract/pdftoppm not installed")

_FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"


def _table_png(lines: list[str]) -> bytes:
    from PIL import Image, ImageDraw, ImageFont
    try:
        font = ImageFont.truetype(_FONT, 34)
    except OSError:
        pytest.skip("DejaVuSans not installed")
    img = Image.new("RGB", (1100, 90 + 60 * len(lines)), "white")
    d = ImageDraw.Draw(img)
    for i, line in enumerate(lines):
        d.text((40, 40 + 60 * i), line, fill="black", font=font)
    out = io.BytesIO()
    img.save(out, "PNG")
    return out.getvalue()


ROWS = ["Depot  Pallets  Code", "Gdynia  23  GDY-7731", "Hamburg  10  HAM-2044"]


def _assert_read(text: str, n: int, where: str, kind: str = "picture") -> None:
    assert f"[Figure {n} · {where} · {kind} · read by OCR, not checked]" in text
    assert f"[End of figure {n}]" in text
    assert "GDY-7731" in text and "HAM-2044" in text


def test_image_file_is_one_figure(tmp_path: Path) -> None:
    from finetune_studio.data.parsers import image

    p = tmp_path / "table.png"
    p.write_bytes(_table_png(ROWS))
    out = image.parse(p)
    _assert_read(out["text"], 1, "image")
    assert out["structured"]["figures"][0]["where"] == "image"


def test_docx_picture_sits_between_its_paragraphs(tmp_path: Path) -> None:
    docx = pytest.importorskip("docx")
    from finetune_studio.data.parsers import docx as parser

    d = docx.Document()
    d.add_paragraph("Loading limits per depot:")
    d.add_picture(io.BytesIO(_table_png(ROWS)))
    d.add_paragraph("Pharma loads are limited to 20.")
    icon = io.BytesIO()
    __import__("PIL.Image").Image.new("RGB", (40, 40), "red").save(icon, "PNG")
    d.add_picture(io.BytesIO(icon.getvalue()))  # an icon: never a figure
    p = tmp_path / "sop.docx"
    d.save(str(p))
    text = parser.parse(p)["text"]
    _assert_read(text, 1, "paragraph 2")
    assert text.index("Loading limits") < text.index("[Figure 1") < text.index("Pharma loads")
    assert "[Figure 2" not in text


def test_pptx_picture_lands_on_its_slide(tmp_path: Path) -> None:
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    from finetune_studio.data.parsers import pptx as parser

    prs = pptx.Presentation()
    prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = "Intro"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Depots"
    s2.shapes.add_picture(io.BytesIO(_table_png(ROWS)), Inches(1), Inches(2))
    p = tmp_path / "deck.pptx"
    prs.save(str(p))
    text = parser.parse(p)["text"]
    _assert_read(text, 1, "slide 2")
    assert text.index("=== Slide 2 ===") < text.index("[Figure 1")


def _text_pdf_page(text: str) -> bytes:
    stream = f"BT /F1 14 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
            b"/Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for i, o in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + o + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref))
    return out.getvalue()


def test_pdf_scanned_page_inside_a_digital_pdf_is_read(tmp_path: Path) -> None:
    from PIL import Image
    from pypdf import PdfReader, PdfWriter

    from finetune_studio.data.parsers import pdf as parser

    scan = io.BytesIO()
    Image.open(io.BytesIO(_table_png(ROWS))).convert("RGB").save(scan, "PDF", resolution=150)
    w = PdfWriter()
    for src in (_text_pdf_page("Reefer depot directory, page one with real text on it."), scan.getvalue()):
        for page in PdfReader(io.BytesIO(src)).pages:
            w.add_page(page)
    p = tmp_path / "mixed.pdf"
    with p.open("wb") as f:
        w.write(f)
    out = parser.parse(p)
    assert "page one with real text" in out["text"]
    _assert_read(out["text"], 1, "page 2", "scanned page")
    assert out["structured"]["figures"] == [{"n": 1, "where": "page 2", "whole_page": True,
                                             "bytes": out["structured"]["figures"][0]["bytes"]}]


def test_block_regex_round_trips() -> None:
    fig = figures.Figure(3, "page 4", 4, b"")
    b = figures.block(fig, "a | b\n1 | 2", "OCR, not checked")
    m = figures.BLOCK_RE.search(f"before\n{b}\nafter")
    assert m and m.group(1) == "3" and m.group(3) == "a | b\n1 | 2"


def test_pdf_shared_resources_list_each_chart_once_on_its_page() -> None:
    # every page of this PDF lists both chart images in a shared resource dictionary; only pages 3 and 5 draw them
    p = next(Path("tests/corpus/korvane/files").rglob("annual_report_2023_extract.pdf"))
    assert [f.where for f in figures.extract(p)] == ["page 3", "page 5"]

